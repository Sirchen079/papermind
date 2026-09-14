"""Editable report export with semantic paragraphs, tables and readable math.

Pagination is deterministic and never replaces overflow with an ellipsis.
"""
import io
import re
import unicodedata
from dataclasses import dataclass
from datetime import datetime

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from app.models import Report

MAX_LINES = 10
_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_LIST_RE = re.compile(r"^\s*(?:[-*+]\s+|\d+[.)]\s+)")


@dataclass
class Block:
    kind: str
    text: str = ""
    rows: list[list[str]] | None = None


def _math_text(value: str) -> str:
    """Readable, editable notation for common inline LaTeX, without changing numbers."""
    value = re.sub(r"\\(?:text|mathrm|mathbf|operatorname)\{([^{}]*)\}", r"\1", value)
    for _ in range(8):
        changed = re.sub(r"\\(?:d?frac)\{([^{}]*)\}\{([^{}]*)\}", r"(\1)/(\2)", value)
        if changed == value:
            break
        value = changed
    commands = {"alpha": "α", "beta": "β", "gamma": "γ", "delta": "δ", "epsilon": "ε",
                "theta": "θ", "lambda": "λ", "mu": "μ", "sigma": "σ", "pi": "π", "omega": "ω",
                "Delta": "Δ", "Sigma": "Σ", "sum": "∑", "prod": "∏", "infty": "∞",
                "times": "×", "cdot": "·", "leq": "≤", "geq": "≥", "neq": "≠", "approx": "≈",
                "pm": "±", "rightarrow": "→", "in": "∈", "log": "log", "exp": "exp"}
    value = re.sub(r"\\([A-Za-z]+)", lambda m: commands.get(m[1], m[0]), value)
    value = re.sub(r"\\sqrt\{([^{}]*)\}", r"√(\1)", value)
    value = value.replace(r"\left", "").replace(r"\right", "").replace(r"\,", " ").replace(r"\!", "")
    subs = dict(zip("0123456789+-=()aehijklmnoprstuvx", "₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎ₐₑₕᵢⱼₖₗₘₙₒₚᵣₛₜᵤᵥₓ"))
    supers = dict(zip("0123456789+-=()in", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾ⁱⁿ"))
    def script(match):
        marker, braced, single = match.groups()
        body = braced if braced is not None else single
        mapping = subs if marker == "_" else supers
        return "".join(mapping[c] for c in body) if all(c in mapping for c in body) else marker + "(" + body + ")"
    return re.sub(r"([_^])(?:\{([^{}]+)\}|([A-Za-z0-9]))", script, value)


def _clean_text(line: str) -> str:
    line = re.sub(r"(?<!\\)\$\$?(.+?)(?<!\\)\$\$?", lambda m: _math_text(m[1]), line)
    line = re.sub(r"\\\((.+?)\\\)", lambda m: _math_text(m[1]), line)
    line = re.sub(r"\*\*(.+?)\*\*|`(.+?)`", lambda m: m[1] or m[2] or "", line)
    return line.strip()


def _cells(line: str) -> list[str]:
    return [_clean_text(cell.replace(r"\|", "|")) for cell in re.split(r"(?<!\\)\|", line.strip().strip("|"))]


def _table_separator(line: str) -> bool:
    return "|" in line and all(re.fullmatch(r":?-{3,}:?", cell.strip()) for cell in line.strip().strip("|").split("|"))


def _parse_sections(markdown: str):
    title, subtitle, sections = "组会汇报", "", []
    current = None
    lines = markdown.splitlines()
    i = 0
    while i < len(lines):
        raw = lines[i]
        i += 1
        if not raw.strip():
            continue
        fence = re.match(r"^\s*(`{3,}|~{3,})", raw)
        if fence:
            code = []
            while i < len(lines) and not lines[i].lstrip().startswith(fence[1]):
                code.append(lines[i]); i += 1
            i += int(i < len(lines))
            if current is None:
                current = ("内容", []); sections.append(current)
            current[1].extend(Block("code", line) for line in code)
            continue
        heading = _HEADING_RE.match(raw.strip())
        if heading:
            level, text = len(heading[1]), _clean_text(heading[2])
            if level == 1 and not sections:
                title = text; continue
            if level <= 2:
                current = (text, []); sections.append(current); continue
            if current is None:
                current = ("内容", []); sections.append(current)
            current[1].append(Block("heading", text)); continue
        if i < len(lines) and "|" in raw and _table_separator(lines[i]):
            rows = [_cells(raw)]; i += 1
            while i < len(lines) and lines[i].strip() and "|" in lines[i]:
                rows.append(_cells(lines[i])); i += 1
            if current is None:
                current = ("内容", []); sections.append(current)
            current[1].append(Block("table", rows=rows)); continue
        if raw.strip() in {"$$", r"\["}:
            closer = "$$" if raw.strip() == "$$" else r"\]"
            formula = []
            while i < len(lines) and lines[i].strip() != closer:
                formula.append(lines[i]); i += 1
            i += int(i < len(lines))
            if current is None:
                current = ("内容", []); sections.append(current)
            current[1].append(Block("paragraph", _math_text(" ".join(formula)))); continue
        listed = bool(_LIST_RE.match(raw))
        text = _clean_text(_LIST_RE.sub("", raw))
        if current is None:
            subtitle = (subtitle + " " + text).strip()
        else:
            current[1].append(Block("bullet" if listed else "paragraph", text))
    return title, subtitle, sections


def _wrap_line(line: str, width: int = 64) -> list[str]:
    """Measure conservative CJK width, preferring word boundaries for Latin text."""
    rows, current, units = [], "", 0
    for char in line:
        size = 2 if unicodedata.east_asian_width(char) in ("W", "F") else 1
        if current and units + size > width:
            boundary = max(current.rfind(" "), current.rfind("，"), current.rfind("。"), current.rfind("；")) + 1
            if boundary > len(current) // 2:
                rows.append(current[:boundary]); current = current[boundary:]
                units = sum(2 if unicodedata.east_asian_width(c) in ("W", "F") else 1 for c in current)
            else:
                rows.append(current); current = ""; units = 0
        current += char; units += size
    if current:
        rows.append(current)
    return rows or [""]


def _paragraph_style(paragraph, kind: str):
    paragraph.font.name = "Consolas" if kind == "code" else "Microsoft YaHei"
    paragraph.font.size = Pt(18)
    paragraph.font.bold = kind == "heading"
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(6)
    paragraph.line_spacing = 1.15
    properties = paragraph._p.get_or_add_pPr()
    for child in list(properties):
        if child.tag.rsplit("}", 1)[-1] in {"buNone", "buChar", "buAutoNum"}:
            properties.remove(child)
    bullet = OxmlElement("a:buChar" if kind == "bullet" else "a:buNone")
    if kind == "bullet":
        bullet.set("char", "•")
    else:
        properties.set("marL", "0")
        properties.set("indent", "0")
    properties.append(bullet)


def report_pptx_bytes(report: Report) -> bytes:
    title, preamble, sections = _parse_sections(report.content)
    subtitle = _default_subtitle(report)
    if preamble:
        if len(_wrap_line(preamble)) > 2:
            sections.insert(0, ("概述", [Block("paragraph", preamble)]))
        else:
            subtitle = f"{subtitle} · {preamble}"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle

    for section_title, blocks in sections:
        page = 0
        pending: list[Block] = []
        used = 0
        def new_slide():
            nonlocal page
            page += 1
            result = prs.slides.add_slide(prs.slide_layouts[1])
            result.shapes.title.text = section_title + (f"（续 {page}）" if page > 1 else "")
            return result
        def flush():
            nonlocal pending, used
            if not pending:
                return
            result = new_slide()
            body = result.placeholders[1].text_frame
            body.word_wrap = True
            for index, block in enumerate(pending):
                paragraph = body.paragraphs[0] if index == 0 else body.add_paragraph()
                paragraph.text = block.text
                _paragraph_style(paragraph, block.kind)
            pending, used = [], 0
        for block in blocks or [Block("paragraph", "（本节无内容）")]:
            if block.kind == "table":
                flush()
                rows = block.rows or [[]]
                columns = max(len(row) for row in rows)
                # Wide tables keep every column; place groups on separate slides.
                for col_start in range(0, columns, 4):
                    group = list(range(col_start, min(col_start + 4, columns)))
                    selected = [[row[c] if c < len(row) else "" for c in group] for row in rows]
                    width = max(12, 60 // len(group))
                    header = selected[0]
                    header_lines = max(len(_wrap_line(cell, width)) for cell in header)
                    body_budget = max(2, MAX_LINES - min(header_lines, 4))
                    chunks: list[list[list[str]]] = []
                    current_rows: list[list[str]] = []
                    count = 0
                    for row in selected[1:]:
                        wrapped = [_wrap_line(cell, width) for cell in row]
                        row_lines = max(map(len, wrapped))
                        for offset in range(0, row_lines, body_budget):
                            height = min(body_budget, row_lines - offset)
                            if current_rows and count + height > body_budget:
                                chunks.append(current_rows); current_rows, count = [], 0
                            current_rows.append(["".join(parts[offset:offset + body_budget]) for parts in wrapped])
                            count += height
                    if current_rows or not chunks:
                        chunks.append(current_rows)
                    for chunk in chunks:
                        result = new_slide()
                        body_placeholder = result.placeholders[1]
                        shape = result.shapes.add_table(len(chunk) + 1, len(group), body_placeholder.left, body_placeholder.top,
                                                       body_placeholder.width, Inches(4.6))
                        table = shape.table
                        table.first_row = True
                        display_rows = [header, *chunk]
                        measures = [max(len(_wrap_line(cell, width)) for cell in row) for row in display_rows]
                        for row_index, row in enumerate(display_rows):
                            table.rows[row_index].height = Inches(0.2 + measures[row_index] * 0.3)
                            for column_index, value in enumerate(row):
                                cell = table.cell(row_index, column_index)
                                cell.text = value
                                cell.margin_left = cell.margin_right = Inches(0.10)
                                cell.margin_top = cell.margin_bottom = Inches(0.06)
                                cell.fill.solid()
                                cell.fill.fore_color.rgb = RGBColor.from_string("EAE6DF" if row_index == 0 else "FAF9F6")
                                for paragraph in cell.text_frame.paragraphs:
                                    _paragraph_style(paragraph, "heading" if row_index == 0 else "paragraph")
                                    paragraph.font.size = Pt(17)
                                    paragraph.font.color.rgb = RGBColor.from_string("282724")
                                    paragraph.space_after = Pt(0)
                        # Keep an empty content placeholder for editor interoperability.
                        body_placeholder.text = ""
                continue
            wrapped = _wrap_line(block.text)
            if len(wrapped) <= MAX_LINES:
                if used + len(wrapped) > MAX_LINES:
                    flush()
                pending.append(block); used += len(wrapped)
            else:
                flush()
                for start in range(0, len(wrapped), MAX_LINES):
                    pending = [Block(block.kind if start == 0 or block.kind != "bullet" else "paragraph", "".join(wrapped[start:start + MAX_LINES]))]
                    flush()
        flush()
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def _default_subtitle(report: Report) -> str:
    return f"{_fmt(report.since)} 至 {_fmt(report.until)}"


def _fmt(value: datetime | None) -> str:
    return value.strftime("%Y-%m-%d") if value else ""

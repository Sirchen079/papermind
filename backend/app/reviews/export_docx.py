"""Review prose as Word (.docx) (P5b).

`references.render_styled_markdown` first renumbers the body `[Pn]` marks by
first appearance and rebuilds the citation section; this module then converts
the resulting markdown line by line with python-docx, the same way the map
export does. Only the markdown subset the prose actually uses is handled
(headings / bullets / quotes / bold spans); tables, code blocks and images
stay plain text on purpose.
"""
import re
from io import BytesIO

from app.reviews import references

_HEADING = re.compile(r'^(#{1,6})\s+(.*)$')


def _fill(paragraph, text: str) -> None:
    """行内 **加粗** 拆成多个 run（bold 开关），其余行内语法原样保留。"""
    for index, chunk in enumerate(re.split(r'\*\*', text)):
        if chunk:
            paragraph.add_run(chunk).bold = index % 2 == 1


def render_review_docx(session, review_id, style: str = 'gbt') -> bytes:
    from docx import Document

    from app.reviews import export_map, service

    row = service.get(session, review_id)
    cited = references.papers_for(session, references.first_appearance_ids(row.content))
    md = references.render_styled_markdown(row.content, cited, style)
    md = export_map._clean_text(md)  # PDF 提取的文本可能含 XML 不允许的控制字符
    doc = Document()
    titled = False
    for raw in md.splitlines():
        line = raw.strip()
        if not line or line == '---':
            continue
        if line.startswith('>'):
            _fill(doc.add_paragraph(), line.lstrip('>').strip())
            continue
        heading = _HEADING.match(line)
        if heading:
            level = min(len(heading.group(1)), 3)
            if level == 1 and not titled:
                level, titled = 0, True
            _fill(doc.add_heading('', level), heading.group(2).strip())
            continue
        if line.startswith(('- ', '* ')):
            _fill(doc.add_paragraph(style='List Bullet'), line[2:].strip())
            continue
        _fill(doc.add_paragraph(), line)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()

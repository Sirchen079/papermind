"""Lifecycle reproductions for the final release acceptance findings."""
from datetime import datetime
from io import BytesIO
from unittest.mock import patch

import pytest
from pptx import Presentation
from sqlmodel import Session, select

from app.db.engine import get_engine
from app.models import Message, Paper, Provider, Report
from app.reports.pptx import report_pptx_bytes


def _context():
    return object(), Provider(id=1, name="qa", base_url="http://127.0.0.1", api_key_enc=""), "qa-model"


@pytest.mark.parametrize("endpoint", ["messages", "messages/stream"])
def test_failed_chat_survives_reload_and_retry_reuses_question_and_context(client, endpoint):
    cid = client.post('/api/chat/conversations').json()['id']
    calls = []
    def failed(*args, **kwargs):
        calls.append(args[3])
        yield 'error', {'message': '503 temporarily unavailable'}
    def success(*args, **kwargs):
        calls.append(args[3])
        yield 'done', {'content': 'recovered answer', 'tokens': 2}
    path = f'/api/chat/conversations/{cid}/{endpoint}'
    with patch('app.api.chat_api.pick_llm', return_value=_context()), patch('app.agent.loop.run_agent', side_effect=failed):
        result = client.post(path, json={'content': 'Compare the baselines', 'selected_text': 'Original selection'})
    assert result.status_code in (200, 500)
    rows = client.get(f'/api/chat/conversations/{cid}').json()['messages']
    assert len(rows) == 1
    assert rows[0]['delivery_status'] == 'failed'
    assert '503' in rows[0]['error_message']
    assert rows[0]['retryable'] is True
    with patch('app.api.chat_api.pick_llm', return_value=_context()), patch('app.agent.loop.run_agent', side_effect=success):
        retry = client.post(path, json={'content': 'Compare the baselines', 'retry_message_id': rows[0]['id'], 'selected_text': 'WRONG new selection'})
        assert retry.status_code == 200
        stale = client.post(path, json={'content': 'Compare the baselines', 'retry_message_id': rows[0]['id']})
    assert stale.status_code == 409
    assert calls[0] == calls[1]
    assert 'Original selection' in str(calls[1])
    assert 'WRONG new selection' not in str(calls[1])
    finished = client.get(f'/api/chat/conversations/{cid}').json()['messages']
    assert [m['role'] for m in finished] == ['user', 'assistant']
    assert finished[0]['delivery_status'] == 'complete'
    assert finished[1]['content'] == 'recovered answer'


def test_pending_turn_after_restart_can_retry_and_new_question_invalidates_old_retry(client):
    cid = client.post('/api/chat/conversations').json()['id']
    with Session(get_engine()) as session:
        row = Message(conversation_id=cid, role='user', content='Interrupted', delivery_status='pending')
        session.add(row); session.commit(); session.refresh(row); mid = row.id
    assert client.get(f'/api/chat/conversations/{cid}').json()['messages'][0]['retryable']
    with Session(get_engine()) as session:
        session.add(Message(conversation_id=cid, role='user', content='Newer', delivery_status='failed'))
        session.commit()
    with patch('app.api.chat_api.pick_llm', return_value=_context()):
        response = client.post(f'/api/chat/conversations/{cid}/messages/stream', json={'content': 'Interrupted', 'retry_message_id': mid})
    assert response.status_code == 409


def test_pdf_page_activity_updates_recent_reading_without_overwriting_status(client):
    pid = client.post('/api/papers/manual', json={'title': 'Recently read PDF'}).json()['id']
    from app.reading.service import patch_reading_state
    with Session(get_engine()) as session:
        before = patch_reading_state(session, pid, {'status': 'read'})
        after = patch_reading_state(session, pid, {'last_page': 2})
        assert after['last_page'] == 2
        assert after['last_read_at'] >= before['last_read_at']
        assert after['status'] == 'read'
        cleared = patch_reading_state(session, pid, {'last_page': None})
        assert cleared['last_read_at'] == after['last_read_at']
    other = client.post('/api/papers/manual', json={'title': 'No manual status'}).json()['id']
    with Session(get_engine()) as session:
        state = patch_reading_state(session, other, {'last_page': 1})
        assert state['last_read_at'] is not None
        assert state['started_at'] is not None


def test_pptx_tables_math_paragraphs_and_code_are_readable_and_editable():
    report = Report(since=datetime(2026, 9, 1), until=datetime(2026, 9, 10), content='''# Test
## Findings
This is one complete paragraph, not a list of separate lines.

| Method | Score |
| --- | --- |
| Sparse | 0.7 |
| Dense | 0.8 |

## Next
- Metric: $F_1=2PR/(P+R)$
```python
# Keep the comment as code
print("end")
```
''')
    deck = Presentation(BytesIO(report_pptx_bytes(report)))
    tables = [shape.table for slide in deck.slides for shape in slide.shapes if shape.has_table]
    assert len(tables) == 1
    assert [[cell.text for cell in row.cells] for row in tables[0].rows] == [['Method', 'Score'], ['Sparse', '0.7'], ['Dense', '0.8']]
    paragraphs = [p for slide in deck.slides for shape in slide.shapes if shape.has_text_frame for p in shape.text_frame.paragraphs]
    assert any('F₁=2PR/(P+R)' in p.text and p._p.xpath('./a:pPr/a:buChar') for p in paragraphs)
    prose = next(p for p in paragraphs if p.text.startswith('This is one'))
    assert prose._p.xpath('./a:pPr/a:buNone')
    assert any(p.text == '# Keep the comment as code' for p in paragraphs)
    assert all('| ---' not in p.text for p in paragraphs)

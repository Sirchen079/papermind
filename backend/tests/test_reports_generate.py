"""P14.2 group-meeting report generation — mocked LLM, bundled template skill."""

from unittest.mock import MagicMock

from sqlmodel import Session

from app.db.engine import get_engine
from app.models import Report, Skill
from app.reports.service import AGGREGATES_TOKEN, TEMPLATE_SKILL_NAME

REPORT_MD = """# 组会汇报（2026-06-04 至 2026-06-10）

## 本周进展
新入库 1 篇，读完 1 篇。

## 文献收获
- 窗内论文A

## 实验进展
新建 1 个实验，追加日志 1 条。

## 问题与求助
基线复现失败，求助显存不足怎么办。

## 下周计划
- 继续实验
"""


def _seed_template_from_file(client) -> str:
    import re
    from pathlib import Path

    text = Path("user_skills/group-meeting-report.md").read_text(encoding="utf-8")
    parts = text.split("---", 2)
    body = parts[2].strip() if text.lstrip().startswith("---") else text.strip()
    with Session(get_engine()) as session:
        session.add(
            Skill(
                name=TEMPLATE_SKILL_NAME,
                description="组会周报生成模板",
                type="template",
                trigger="manual",
                body=body,
            )
        )
        session.commit()
    return body


def test_bundled_template_skill_is_well_formed():
    """The bundled skill file must parse via the existing skills loader with
    the right type/trigger and carry the aggregates token."""
    import json
    from pathlib import Path

    from app.skills.loader import parse_skill_file

    path = Path("user_skills/group-meeting-report.md")
    data = parse_skill_file(path)
    assert data["name"] == TEMPLATE_SKILL_NAME
    assert data["type"] == "template"
    assert data["trigger"] == "manual"
    assert AGGREGATES_TOKEN in data["body"]
    keywords = json.loads(data["keywords_json"])
    assert keywords  # 组会/周报/汇报


def test_generate_report_success_stores_report(client, monkeypatch):
    body = _seed_template_from_file(client)
    prompts: list[str] = []

    def fake_pick(session, role):
        assert role == "chat"

        class _C:
            def complete(self, provider, model_id, messages, request_kind=None):
                prompts.append(messages[0]["content"])
                return MagicMock(content=REPORT_MD)

        return _C(), object(), "mock-model"

    monkeypatch.setattr("app.providers.selection.pick_llm", fake_pick)

    res = client.post(
        "/api/reports/weekly/generate",
        json={"since": "2026-06-04", "until": "2026-06-10", "problems": "基线复现失败"},
    )
    assert res.status_code == 201, res.text
    report = res.json()
    assert report["model"] == "mock-model"
    assert "## 本周进展" in report["content"]
    assert report["since"].startswith("2026-06-04")
    assert report["until"].startswith("2026-06-10")

    # Prompt carried the template instructions + deterministic aggregates + problems.
    assert "组会汇报" in prompts[0]
    assert AGGREGATES_TOKEN not in prompts[0]  # token was replaced
    assert "统计区间：2026-06-04 至 2026-06-10" in prompts[0]
    assert "基线复现失败" in prompts[0]

    # History list (newest first) + detail.
    listed = client.get("/api/reports").json()
    assert [r["id"] for r in listed] == [report["id"]]
    assert client.get(f"/api/reports/{report['id']}").json()["content"] == REPORT_MD.strip()
    assert client.get("/api/reports/99999").status_code == 404

    with Session(get_engine()) as session:
        row = session.get(Report, report["id"])
        assert row is not None and row.content == REPORT_MD.strip()


def test_generate_report_400_when_no_llm(client):
    _seed_template_from_file(client)
    res = client.post("/api/reports/weekly/generate", json={})
    assert res.status_code == 400
    assert "LLM" in res.json()["detail"]
    assert client.get("/api/reports").json() == []  # nothing stored on failure


def test_generate_report_400_when_template_missing_or_broken(client, monkeypatch, tmp_path):
    # Template nowhere: no DB row and the bundled skills dir replaced by an
    # empty one (the service would otherwise lazily load the bundled file).
    monkeypatch.setattr("app.api.skills_api.default_skills_dir", lambda: tmp_path)
    res = client.post("/api/reports/weekly/generate", json={})
    assert res.status_code == 400
    assert TEMPLATE_SKILL_NAME in res.json()["detail"]

    # Template present but token removed by an editing mistake (the DB row
    # wins over the loader, so the broken body is what generation sees).
    with Session(get_engine()) as session:
        session.add(
            Skill(name=TEMPLATE_SKILL_NAME, type="template", body="写一份汇报。", trigger="manual")
        )
        session.commit()

    def fake_pick(session, role):
        class _C:
            def complete(self, *args, **kwargs):
                return MagicMock(content="x")

        return _C(), object(), "m"

    monkeypatch.setattr("app.providers.selection.pick_llm", fake_pick)
    res = client.post("/api/reports/weekly/generate", json={})
    assert res.status_code == 400
    assert AGGREGATES_TOKEN in res.json()["detail"]


def test_generate_report_502_on_llm_failure(client, monkeypatch):
    _seed_template_from_file(client)

    def fake_pick(session, role):
        class _Boom:
            def complete(self, *args, **kwargs):
                raise RuntimeError("upstream 500")

        return _Boom(), object(), "m"

    monkeypatch.setattr("app.providers.selection.pick_llm", fake_pick)
    res = client.post("/api/reports/weekly/generate", json={})
    assert res.status_code == 502
    assert "upstream 500" in res.json()["detail"]
    assert client.get("/api/reports").json() == []

    # Empty model output is also a 502.
    def fake_empty(session, role):
        class _C:
            def complete(self, *args, **kwargs):
                return MagicMock(content="")

        return _C(), object(), "m"

    monkeypatch.setattr("app.providers.selection.pick_llm", fake_empty)
    assert client.post("/api/reports/weekly/generate", json={}).status_code == 502


def test_generate_report_uses_edited_template_from_db(client, monkeypatch):
    """The DB row wins: users can customize the template in the Skills page."""
    with Session(get_engine()) as session:
        session.add(
            Skill(
                name=TEMPLATE_SKILL_NAME,
                type="template",
                trigger="manual",
                body="自定义指令。请总结：\n{{aggregates}}",
            )
        )
        session.commit()

    prompts: list[str] = []

    def fake_pick(session, role):
        class _C:
            def complete(self, provider, model_id, messages, request_kind=None):
                prompts.append(messages[0]["content"])
                return MagicMock(content="# 自定义汇报")

        return _C(), object(), "m"

    monkeypatch.setattr("app.providers.selection.pick_llm", fake_pick)
    res = client.post("/api/reports/weekly/generate", json={})
    assert res.status_code == 201
    assert prompts[0].startswith("自定义指令。请总结：")
    assert "统计区间" in prompts[0]

    # Sanity: the aggregates endpoint itself stayed LLM-free.
    agg = client.get("/api/reports/weekly").json()
    assert "papers_new" in agg


def test_acceptance_aggregate_generate_download_end_to_end(client, monkeypatch):
    """P14 验收演示（代码级）：过去一周有入库/阅读/实验 → 预览聚合 → mock 生成
    → 下载 markdown 与 pptx（python-pptx 重新打开解析页数与标题）。"""
    from datetime import datetime
    from io import BytesIO

    from pptx import Presentation

    from app.models import Experiment, ExperimentLog, Paper, PaperReadingState, Project

    inside = datetime(2026, 6, 8, 10, 0, 0)
    with Session(get_engine()) as session:
        paper = Paper(source="manual", title="验收论文", created_at=inside)
        session.add(paper)
        session.commit()
        session.refresh(paper)
        session.add(PaperReadingState(paper_id=paper.id, status="read", finished_at=inside))
        project = Project(name="验收项目")
        session.add(project)
        session.commit()
        session.refresh(project)
        experiment = Experiment(
            name="验收实验", project_id=project.id, status="running", created_at=inside
        )
        session.add(experiment)
        session.commit()
        session.refresh(experiment)
        session.add(ExperimentLog(experiment_id=experiment.id, content="验收日志", created_at=inside))
        session.commit()

    # 1) 预览聚合数据（确定性）。
    agg = client.get("/api/reports/weekly", params={"since": "2026-06-04", "until": "2026-06-10"}).json()
    assert agg["papers_new"]["count"] == 1
    assert agg["papers_read"]["count"] == 1
    assert agg["experiments"]["created"]["count"] == 1
    assert agg["experiments"]["logs_added"]["count"] == 1

    # 2) 模板 + mock LLM 生成。
    body = _seed_template_from_file(client)

    def fake_pick(session, role):
        class _C:
            def complete(self, provider, model_id, messages, request_kind=None):
                return MagicMock(content=REPORT_MD)

        return _C(), object(), "mock-model"

    monkeypatch.setattr("app.providers.selection.pick_llm", fake_pick)
    created = client.post(
        "/api/reports/weekly/generate",
        json={"since": "2026-06-04", "until": "2026-06-10", "problems": "求基线调参建议"},
    )
    assert created.status_code == 201, created.text
    report_id = created.json()["id"]
    assert "组会汇报" in body  # template body drove the prompt

    # 3) markdown 下载 = 存储内容。
    md = client.get(f"/api/reports/{report_id}/markdown")
    assert md.status_code == 200
    assert md.text == REPORT_MD.strip()

    # 4) pptx 下载可被 python-pptx 重新打开并解析。
    pptx_res = client.get(f"/api/reports/{report_id}/pptx")
    assert pptx_res.status_code == 200
    prs = Presentation(BytesIO(pptx_res.content))
    slides = list(prs.slides)
    assert len(slides) == 6  # 标题页 + 五节
    assert slides[0].shapes.title.text.startswith("组会汇报")
    assert [s.shapes.title.text for s in slides[1:]] == [
        "本周进展",
        "文献收获",
        "实验进展",
        "问题与求助",
        "下周计划",
    ]

    # 5) 历史列表包含本次生成。
    assert [r["id"] for r in client.get("/api/reports").json()] == [report_id]

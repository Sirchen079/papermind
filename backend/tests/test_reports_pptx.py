"""P14.3 markdown / PPTX export — python-pptx re-parses what it produced."""

from datetime import datetime
from io import BytesIO

from pptx import Presentation
from sqlmodel import Session

from app.db.engine import get_engine
from app.models import Report
from app.reports.pptx import MAX_LINES, report_pptx_bytes

REPORT_MD = """# 组会汇报（2026-06-04 至 2026-06-10）

前言段：本周整体平稳。

## 本周进展
- 新入库 2 篇
- 读完 1 篇

## 文献收获
1. 窗内论文A
2. **窗内论文B** 与 `关键词`

## 实验进展
- 新建 1 个实验

## 问题与求助
本周暂无

## 下周计划
- 继续实验
"""


def _seed_report(content: str = REPORT_MD) -> int:
    with Session(get_engine()) as session:
        row = Report(
            since=datetime(2026, 6, 4),
            until=datetime(2026, 6, 10),
            content=content,
            model="mock-model",
        )
        session.add(row)
        session.commit()
        session.refresh(row)
        return row.id


def test_pptx_reopens_with_expected_slides_and_titles(client):
    report_id = _seed_report()
    res = client.get(f"/api/reports/{report_id}/pptx")
    assert res.status_code == 200, res.text
    assert res.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )
    assert "attachment" in res.headers["content-disposition"]

    prs = Presentation(BytesIO(res.content))
    # Title page + one slide per H2 section (5 sections).
    slides = list(prs.slides)
    assert len(slides) == 6
    title_slide = slides[0]
    assert title_slide.shapes.title.text == "组会汇报（2026-06-04 至 2026-06-10）"
    subtitle = title_slide.placeholders[1].text
    assert "2026-06-04" in subtitle and "2026-06-10" in subtitle
    assert "前言段" in subtitle

    expected_titles = ["本周进展", "文献收获", "实验进展", "问题与求助", "下周计划"]
    for slide, expected in zip(slides[1:], expected_titles):
        assert slide.shapes.title.text == expected

    # List markers and emphasis stripped; content lines carried over.
    body = prs.slides[1].placeholders[1].text_frame
    texts = [p.text for p in body.paragraphs]
    assert texts == ["新入库 2 篇", "读完 1 篇"]
    lit = prs.slides[2].placeholders[1].text_frame
    lit_texts = [p.text for p in lit.paragraphs]
    assert lit_texts == ["窗内论文A", "窗内论文B 与 关键词"]


def test_markdown_download_returns_stored_content(client):
    report_id = _seed_report()
    res = client.get(f"/api/reports/{report_id}/markdown")
    assert res.status_code == 200
    assert res.headers["content-type"].startswith("text/markdown")
    assert "## 实验进展" in res.text
    assert "attachment" in res.headers["content-disposition"]


def test_pptx_paginates_long_sections_without_losing_content(client):
    lines = "\n".join(f"- 条目 {i}" for i in range(MAX_LINES + 5))
    markdown = f"# 汇报\n\n## 长清单\n{lines}\n"
    report_id = _seed_report(markdown)

    res = client.get(f"/api/reports/{report_id}/pptx")
    assert res.status_code == 200
    prs = Presentation(BytesIO(res.content))
    assert len(prs.slides) == 3  # title + two section slides
    texts = [p.text for slide in list(prs.slides)[1:] for p in slide.placeholders[1].text_frame.paragraphs]
    assert texts == [f"条目 {i}" for i in range(MAX_LINES + 5)]
    assert '续' in prs.slides[2].shapes.title.text

    # Fully empty content still yields a valid deck (title page only).
    empty_id = _seed_report("")
    res = client.get(f"/api/reports/{empty_id}/pptx")
    assert res.status_code == 200
    prs = Presentation(BytesIO(res.content))
    assert len(prs.slides) == 1

    # Missing report → 404.
    assert client.get("/api/reports/99999/pptx").status_code == 404
    assert client.get("/api/reports/99999/markdown").status_code == 404


def test_pptx_keeps_long_chinese_paragraph_and_code_content():
    paragraph = '统一数据划分后开展公平比较。' * 60
    report = Report(since=datetime(2026,9,1), until=datetime(2026,9,10),
                    content=f'## 研究结果\n{paragraph}\n```python\nprint("保留结尾")\n```')
    prs = Presentation(BytesIO(report_pptx_bytes(report)))
    text = ''.join(p.text for slide in list(prs.slides)[1:] for p in slide.placeholders[1].text_frame.paragraphs)
    assert text == paragraph + 'print("保留结尾")'
    assert len(prs.slides) > 2
